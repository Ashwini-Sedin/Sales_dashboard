import logging
from datetime import datetime
from io import BytesIO
from uuid import UUID

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Template, TemplateError

from app.models.lead import Lead
from app.models.division import Division
from app.models.user import User, UserRole
from app.models.nda_clause import NdaClause
from app.models.generated_document import GeneratedDocument
from app.models.activity_timeline import ActivityTimeline, ActivityEventType
from app.schemas.document import NdaInput
from app.services.documents.base_document_service import BaseDocumentService
from app.services.template_service import get_active_template
from app.services.notification_service import create_in_app_notification
from app.core.socket_manager import socket_manager
from sqlalchemy import select

logger = logging.getLogger(__name__)

NDA_TYPE_LABELS = {
    "unilateral": "UNILATERAL NON-DISCLOSURE AGREEMENT",
    "mutual": "MUTUAL NON-DISCLOSURE AGREEMENT",
    "multilateral": "MULTILATERAL NON-DISCLOSURE AGREEMENT"
}

class NdaService(BaseDocumentService):
    async def _fetch_nda_data(self, db, lead_id: UUID, inputs: NdaInput) -> dict:
        lead = (await db.execute(select(Lead).where(Lead.id == lead_id))).scalar_one()
        division = (await db.execute(select(Division).where(Division.id == lead.division_id))).scalar_one()
        company_signatory = (await db.execute(select(User).where(User.id == inputs.company_signatory_user_id))).scalar_one()
        
        branding = division.branding_config or {}
        primary_color = branding.get('primary_color', '#0D1B2A')
        secondary_color = branding.get('secondary_color', '#00C6D7')
        logo_url = branding.get('logo_url', None)

        primary_rgb = self._hex_to_rgb(primary_color)
        secondary_rgb = self._hex_to_rgb(secondary_color)

        # Fetch clauses
        default_stmt = select(NdaClause).where(NdaClause.is_default == True).order_by(NdaClause.clause_type.asc())
        default_clauses = (await db.execute(default_stmt)).scalars().all()
        
        custom_clauses = []
        if inputs.custom_clause_ids:
            custom_stmt = select(NdaClause).where(NdaClause.id.in_(inputs.custom_clause_ids))
            custom_clauses = (await db.execute(custom_stmt)).scalars().all()

        # Merge clauses: custom overrides default of same type
        clauses_dict = {c.clause_type: c for c in default_clauses}
        for cc in custom_clauses:
            clauses_dict[cc.clause_type] = cc
        
        all_clauses = list(clauses_dict.values())

        template_context = {
            "company": division.name,
            "company_legal": division.name,
            "company_address": "Address on file", # Placeholder as requested
            "client": inputs.client_legal_name,
            "client_address": inputs.client_address,
            "effective_date": inputs.effective_date,
            "purpose": inputs.purpose_of_disclosure,
            "term": str(inputs.term_years),
            "nda_type": inputs.nda_type,
            "governing_law": "applicable jurisdiction"
        }

        return {
            "lead": lead,
            "division": division,
            "company_signatory": company_signatory,
            "primary_color": primary_color,
            "secondary_color": secondary_color,
            "logo_url": logo_url,
            "primary_rgb": primary_rgb,
            "secondary_rgb": secondary_rgb,
            "all_clauses": all_clauses,
            "template_context": template_context
        }

    def _render_clause(self, clause_text: str, context: dict) -> str:
        try:
            return Template(clause_text).render(**context)
        except TemplateError as e:
            logger.warning(f"Failed to render clause template: {e}")
            return clause_text

    def _add_signature_block(self, doc, company_signatory, inputs: NdaInput, additional_parties: list) -> None:
        total_parties = 2 + len(additional_parties)
        cols = 2 if total_parties <= 2 else 3
        rows = (total_parties + cols - 1) // cols
        
        table = doc.add_table(rows=rows, cols=cols)
        table.autofit = True
        
        parties_to_render = [
            {"name": company_signatory.division.name if hasattr(company_signatory, 'division') else "The Company", 
             "signatory": f"{company_signatory.first_name} {company_signatory.last_name}", 
             "title": company_signatory.role.value if hasattr(company_signatory.role, 'value') else str(company_signatory.role)},
            {"name": inputs.client_legal_name, 
             "signatory": inputs.client_signatory_name, 
             "title": inputs.client_signatory_title}
        ]
        
        for p in additional_parties:
            parties_to_render.append({
                "name": p.company,
                "signatory": p.name,
                "title": p.designation
            })

        for idx, party in enumerate(parties_to_render):
            r = idx // cols
            c = idx % cols
            cell = table.rows[r].cells[c]
            p = cell.paragraphs[0]
            
            run = p.add_run("SIGNED FOR AND ON BEHALF OF:\n")
            run.bold = True
            
            run = p.add_run(party['name'] + "\n")
            run.bold = True
            run.font.size = Pt(12)
            
            p.add_run("\n\n")
            p.add_run("Signature: _______________________\n")
            p.add_run(f"Name: {party['signatory']}\n")
            p.add_run(f"Title: {party['title']}\n")
            p.add_run("Date: _______________________")

    async def generate_docx(self, db, lead_id: UUID, user_id: UUID, inputs: NdaInput) -> GeneratedDocument:
        # Step 1: Fetch data
        data = await self._fetch_nda_data(db, lead_id, inputs)
        
        # Step 2: Template check
        template_bytes = await get_active_template(db, data['division'].id, 'nda', 'docx')
        if template_bytes:
            doc = Document(BytesIO(template_bytes))
        else:
            doc = Document()

        # Step 3: Margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

        # Step 4: Build NDA Document
        # Cover Section
        if data['logo_url']:
            # Placeholder for logo as we can't fetch URL here easily
            pass
            
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(NDA_TYPE_LABELS[inputs.nda_type])
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = data['primary_rgb']

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text = f"This {NDA_TYPE_LABELS[inputs.nda_type]} (\"Agreement\") is entered into as of {inputs.effective_date} by and between:"
        run = p.add_run(text)
        run.font.size = Pt(11)

        # Parties Block
        def add_party_para(name, label):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = data['primary_rgb']
            p.add_run(f" ({label})")

        label_a = "Disclosing Party" if inputs.nda_type == "unilateral" else "Party A"
        label_b = "Receiving Party" if inputs.nda_type == "unilateral" else "Party B"
        
        add_party_para(data['division'].name, label_a)
        p = doc.add_paragraph("AND")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_party_para(inputs.client_legal_name, label_b)
        
        if inputs.nda_type == "multilateral":
            for idx, party in enumerate(inputs.additional_parties):
                p = doc.add_paragraph("AND")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_party_para(party.company, f"Party {chr(67 + idx)}") # Party C, D...

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"For the purpose of: {inputs.purpose_of_disclosure}")
        run.font.size = Pt(11)
        run.italic = True

        doc.add_page_break()

        # Clause Sections
        CLAUSE_ORDER = [
            "confidentiality", "exclusions", "obligations",
            "term", "governing_law", "dispute_resolution",
            "ip", "entire_agreement", "severability", "waiver"
        ]
        
        section_number = 1
        clauses_by_type = {c.clause_type: c for c in data['all_clauses']}
        
        for c_type in CLAUSE_ORDER:
            clause = clauses_by_type.get(c_type)
            if not clause:
                continue
                
            heading_text = f"{section_number}. {c_type.replace('_', ' ').title()}"
            h = doc.add_heading(heading_text, level=2)
            h.runs[0].font.color.rgb = data['primary_rgb']
            
            rendered_text = self._render_clause(clause.clause_text, data['template_context'])
            p = doc.add_paragraph(rendered_text)
            p.style.font.size = Pt(11)
            
            section_number += 1

        doc.add_page_break()

        # Signature Block
        h = doc.add_heading("IN WITNESS WHEREOF", level=2)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph("The parties have executed this Agreement as of the date first written above.")
        p.style.font.size = Pt(11)
        
        self._add_signature_block(doc, data['company_signatory'], inputs, inputs.additional_parties)

        # Step 5: Save to BytesIO
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        file_bytes = output.read()

        # Step 6: Metadata
        await self.archive_previous_versions(db, lead_id, 'nda')
        version = await self.get_next_version(db, lead_id, 'nda')
        title = self.build_document_title('nda', version, data['lead'].company_name)
        filename = f"{title}.docx"

        # Step 7: S3
        s3_key = await self.upload_to_s3(
            file_bytes, str(lead_id), str(data['division'].id),
            'nda', filename,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        # Step 8: SharePoint
        folder = await self.create_sharepoint_folder(str(lead_id), data['division'].name, data['lead'].company_name, 'nda')
        sharepoint_url = await self.upload_to_sharepoint(file_bytes, folder, filename)

        # Step 9: DB Record
        document = await self.save_document_record(
            db=db, lead_id=lead_id, division_id=data['division'].id,
            template_id=None, doc_type='nda', format='docx',
            version_number=version, title=title, s3_key=s3_key,
            sharepoint_url=sharepoint_url, created_by=user_id,
            status='pending_legal_review'
        )

        # Step 10: Notify Legal
        legal_users = (await db.execute(
            select(User).where(User.role == UserRole.legal, User.division_id == data['division'].id)
        )).scalars().all()
        
        for legal_user in legal_users:
            await create_in_app_notification(
                db=db, user_id=legal_user.id, lead_id=lead_id,
                type='doc_generated',
                message=f"NDA v{version} for {data['lead'].company_name} requires legal review."
            )

        # Step 11: Activity
        activity = ActivityTimeline(
            lead_id=lead_id, actor_id=user_id,
            event_type=ActivityEventType.document_generated,
            description=f"📜 NDA generated (v{version}) — Pending Legal Review",
            metadata_={
                "doc_type": "nda", "format": "docx", "version": version,
                "nda_type": inputs.nda_type, "document_id": str(document.id),
                "status": "pending_legal_review"
            }
        )
        db.add(activity)
        await db.commit()

        # Step 12: Socket
        await socket_manager.emit_to_lead_room(
            lead_id=str(lead_id),
            event="document_generated",
            data={
                "document_id": str(document.id), "doc_type": "nda",
                "format": "docx", "version": version, "title": title,
                "status": "pending_legal_review"
            }
        )

        return document

    def _hex_to_rgb(self, hex_color):
        hex_color = hex_color.lstrip('#')
        return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))
