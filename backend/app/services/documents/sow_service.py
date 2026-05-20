import logging
from datetime import datetime
from io import BytesIO
from uuid import UUID

import requests
import boto3
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlalchemy import select

from app.core.config import settings
from app.core.socket_manager import socket_manager
from app.models.lead import Lead
from app.models.division import Division
from app.models.user import User
from app.models.generated_document import GeneratedDocument
from app.models.activity_timeline import ActivityTimeline, ActivityEventType
from app.schemas.document import SowInput
from app.services.documents.base_document_service import BaseDocumentService
from app.services.template_service import get_active_template
from app.services.notification_service import create_in_app_notification
from app.services.documents.quick_sales_docx import _apply_cell_color, _hex_to_rgb

logger = logging.getLogger(__name__)

METHODOLOGY_LABELS = {
  "agile": "Agile / Scrum",
  "waterfall": "Waterfall",
  "hybrid": "Hybrid (Agile + Waterfall)"
}

IP_OWNERSHIP_LABELS = {
  "client": "Client owns all IP upon full payment",
  "company": "Service Provider retains all IP",
  "joint": "Joint ownership — both parties share IP rights"
}

class SowService(BaseDocumentService):

    async def _fetch_sow_data(
        self, db, lead_id: UUID, inputs: SowInput
    ) -> dict:
        """
        Fetch all data needed for SOW generation.
        """
        lead_stmt = select(Lead).where(Lead.id == lead_id)
        result = await db.execute(lead_stmt)
        lead = result.scalar_one()

        division_stmt = select(Division).where(Division.id == lead.division_id)
        result = await db.execute(division_stmt)
        division = result.scalar_one()

        division_head_stmt = select(User).where(User.id == division.head_user_id)
        result = await db.execute(division_head_stmt)
        division_head = result.scalar_one()

        branding = division.branding_config or {}
        primary_color = branding.get('primary_color', '#0D1B2A')
        secondary_color = branding.get('secondary_color', '#00C6D7')
        logo_url = branding.get('logo_url', None)
        
        primary_rgb = _hex_to_rgb(primary_color)
        secondary_rgb = _hex_to_rgb(secondary_color)

        # Calculate payment totals
        total_contract_value = sum(p.amount for p in inputs.payments)

        # Import from presales if linked
        presales_data = None
        if inputs.linked_presales_doc_id:
            presales_doc_stmt = select(GeneratedDocument).where(
                GeneratedDocument.id == inputs.linked_presales_doc_id
            )
            result = await db.execute(presales_doc_stmt)
            presales_doc = result.scalar_one_or_none()
            
            if presales_doc and presales_doc.doc_type == 'presales':
                presales_data = await self._parse_presales_to_sow(presales_doc)

        return {
            'lead': lead,
            'division': division,
            'division_head': division_head,
            'primary_color': primary_color,
            'secondary_color': secondary_color,
            'logo_url': logo_url,
            'primary_rgb': primary_rgb,
            'secondary_rgb': secondary_rgb,
            'total_contract_value': total_contract_value,
            'presales_data': presales_data
        }

    async def _parse_presales_to_sow(
        self, presales_doc: GeneratedDocument
    ) -> dict | None:
        """
        Extract useful data from a linked presales document
        to pre-populate SOW fields.
        """
        try:
            s3 = boto3.client(
                's3',
                aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
                aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
                region_name=settings.AWS_REGION
            )
            response = s3.get_object(
                Bucket=settings.AWS_S3_DOCUMENTS_BUCKET,
                Key=presales_doc.s3_key
            )
            doc_bytes = response['Body'].read()
            doc = Document(BytesIO(doc_bytes))
            
            scope_items = []
            effort_rows = []
            commercial_model = "TBD"

            for table in doc.tables:
                # Heuristic for scope items (In Scope header)
                if len(table.rows) > 0 and "In Scope" in table.cell(0, 0).text:
                    for row in table.rows[1:]:
                        if row.cells[0].text.strip():
                            scope_items.append(row.cells[0].text.strip().replace("✅ ", ""))
                
                # Heuristic for effort rows
                if len(table.rows) > 0 and "Phase" in table.cell(0, 0).text and "Role" in table.cell(0, 1).text:
                    for row in table.rows[1:]:
                        if row.cells[0].text.strip() != "TOTAL":
                            effort_rows.append({
                                "phase": row.cells[0].text.strip(),
                                "role": row.cells[1].text.strip(),
                                "days": row.cells[2].text.strip()
                            })

            return {
                "scope_items": scope_items,
                "effort_rows": effort_rows,
                "commercial_model": commercial_model
            }
        except Exception as e:
            logger.warning(f"Could not parse presales doc: {e}")
            return None

    async def generate_docx(
        self,
        db,
        lead_id: UUID,
        user_id: UUID,
        inputs: SowInput
    ) -> GeneratedDocument:
        
        # Step 1 — Fetch data
        data = await self._fetch_sow_data(db, lead_id, inputs)
        
        # Step 2 — Check for active template
        template_bytes = await get_active_template(
            db, data['division'].id, 'sow', 'docx'
        )
        if template_bytes:
            doc = Document(BytesIO(template_bytes))
        else:
            doc = Document()

        # Step 3 — Set document margins (legal format)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

        # Step 4 — Build all 15 SOW sections
        
        # Section 1 — Cover Page
        if data['logo_url']:
            try:
                resp = requests.get(data['logo_url'])
                if resp.status_code == 200:
                    doc.add_picture(BytesIO(resp.content), width=Inches(2.0))
                    last_p = doc.paragraphs[-1]
                    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                logger.warning(f"Failed to add logo: {e}")

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("STATEMENT OF WORK")
        run.bold = True
        run.font.size = Pt(28)
        run.font.color.rgb = data['primary_rgb']

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(inputs.project_name)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = data['secondary_rgb']

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Prepared For: {data['lead'].company_name}")
        run.font.size = Pt(14)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Prepared By: {data['division'].name}")
        run.font.size = Pt(14)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Methodology: {METHODOLOGY_LABELS[inputs.methodology]}")
        run.font.size = Pt(12)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Date: {datetime.today().strftime('%B %d, %Y')}")
        run.font.size = Pt(12)
        
        doc.add_page_break()

        # Section 2 — Project Overview
        h = doc.add_heading("1. Project Overview", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        overview_rows = [
            ("Project Name", inputs.project_name),
            ("Client", data['lead'].company_name),
            ("Service Provider", data['division'].name),
            ("Methodology", METHODOLOGY_LABELS[inputs.methodology]),
            ("IP Ownership", IP_OWNERSHIP_LABELS[inputs.ip_ownership]),
            ("Change Control Threshold", f"{inputs.change_control_threshold_days} business days")
        ]
        for label, val in overview_rows:
            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[0].paragraphs[0].runs[0].bold = True
            row_cells[1].text = val
        
        doc.add_page_break()

        # Section 3 — Objectives
        h = doc.add_heading("2. Project Objectives", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        for index, obj in enumerate(inputs.objectives):
            p = doc.add_paragraph(f"{index + 1}. {obj}", style="List Number")
        
        doc.add_page_break()

        # Section 4 — Deliverables
        h = doc.add_heading("3. Deliverables", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_texts = ["Deliverable", "Description", "Acceptance Criteria", "Target Date"]
        for i, text in enumerate(hdr_texts):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _apply_cell_color(hdr_cells[i], data['primary_color'])
        
        for d in inputs.deliverables:
            row_cells = table.add_row().cells
            row_cells[0].text = d.name
            row_cells[1].text = d.description
            row_cells[2].text = d.acceptance_criteria
            row_cells[3].text = d.target_date
        
        doc.add_page_break()

        # Section 5 — Out of Scope
        h = doc.add_heading("4. Out of Scope", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        for item in inputs.out_of_scope:
            doc.add_paragraph("❌ " + item, style="List Bullet")
        
        doc.add_page_break()

        # Section 6 — Methodology
        h = doc.add_heading("5. Methodology", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        p = doc.add_paragraph()
        run = p.add_run(METHODOLOGY_LABELS[inputs.methodology])
        run.bold = True
        run.font.size = Pt(13)
        
        methodology_desc = {
            "agile": "The project will follow an Agile/Scrum methodology with 2-week sprints, daily standups, sprint reviews and retrospectives.",
            "waterfall": "The project will follow a sequential Waterfall methodology with clearly defined phases, gates and sign-offs at each stage.",
            "hybrid": "The project will combine Agile flexibility for development phases with Waterfall structured governance for planning and delivery milestones."
        }
        p = doc.add_paragraph(methodology_desc.get(inputs.methodology, ""))
        p.runs[0].font.size = Pt(11)
        
        doc.add_page_break()

        # Section 7 — Milestones
        h = doc.add_heading("6. Project Milestones", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_texts = ["Milestone", "Due Date", "Payment Linked"]
        for i, text in enumerate(hdr_texts):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _apply_cell_color(hdr_cells[i], data['primary_color'])
        
        for m in inputs.milestones:
            row_cells = table.add_row().cells
            row_cells[0].text = m.name + "\n" + m.description
            row_cells[1].text = m.due_date
            row_cells[2].text = "✅ Yes" if m.payment_linked else "No"
        
        doc.add_page_break()

        # Section 8 — RACI Matrix
        h = doc.add_heading("7. RACI Matrix", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph("R = Responsible | A = Accountable | C = Consulted | I = Informed")
        run = p.runs[0]
        run.font.size = Pt(10)
        run.italic = True
        run.font.color.rgb = RGBColor(128, 128, 128)
        
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_texts = ["Deliverable", "R", "A", "C", "I"]
        for i, text in enumerate(hdr_texts):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _apply_cell_color(hdr_cells[i], data['primary_color'])
        
        for entry in inputs.raci_matrix.entries:
            row_cells = table.add_row().cells
            row_cells[0].text = entry.deliverable
            row_cells[1].text = entry.responsible
            row_cells[2].text = entry.accountable
            row_cells[3].text = entry.consulted
            row_cells[4].text = entry.informed
            
            _apply_cell_color(row_cells[1], "E8F5E9") # Light Green
            _apply_cell_color(row_cells[2], "E3F2FD") # Light Blue
            _apply_cell_color(row_cells[3], "FFFDE7") # Light Yellow
            _apply_cell_color(row_cells[4], "F5F5F5") # Light Gray
        
        doc.add_page_break()

        # Section 9 — Payment Schedule
        h = doc.add_heading("8. Payment Schedule", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph(f"Total Contract Value: ${data['total_contract_value']:,.2f}")
        run = p.runs[0]
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = data['primary_rgb']
        
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_texts = ["Milestone", "Amount", "Due Date"]
        for i, text in enumerate(hdr_texts):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _apply_cell_color(hdr_cells[i], data['primary_color'])
        
        for p_mile in inputs.payments:
            row_cells = table.add_row().cells
            row_cells[0].text = p_mile.milestone_name
            row_cells[1].text = f"${p_mile.amount:,.2f} ({p_mile.percentage:.0f}%)"
            row_cells[2].text = p_mile.due_date
            
        total_row = table.add_row().cells
        total_row[0].text = "TOTAL"
        total_row[1].text = f"${data['total_contract_value']:,.2f}"
        total_row[2].text = ""
        for i in range(2):
            total_row[i].paragraphs[0].runs[0].bold = True
            total_row[i].paragraphs[0].runs[0].font.color.rgb = data['primary_rgb']
            
        doc.add_page_break()

        # Section 10 — IP Ownership
        h = doc.add_heading("9. Intellectual Property", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph(IP_OWNERSHIP_LABELS[inputs.ip_ownership])
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(13)
        
        p = doc.add_paragraph("All work product, deliverables and associated intellectual property created during the course of this engagement shall be owned as specified above upon receipt of full payment. Pre-existing IP of either party remains with the original owner.")
        p.runs[0].font.size = Pt(11)
        
        doc.add_page_break()

        # Section 11 — Change Control
        h = doc.add_heading("10. Change Control Process", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph(f"Any changes to the scope, timeline or budget of this engagement must be submitted as a formal Change Request. Changes will be reviewed and responded to within {inputs.change_control_threshold_days} business days. Approved changes will be documented in a Change Order and signed by both parties before implementation.")
        p.runs[0].font.size = Pt(11)
        
        doc.add_page_break()

        # Section 12 — Assumptions & Dependencies
        h = doc.add_heading("11. Assumptions & Dependencies", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        doc.add_paragraph("This SOW is based on the following assumptions:")
        standard_assumptions = [
            "Client will provide timely access to required systems, data and personnel",
            "Client will designate a project manager as primary point of contact",
            "All feedback and approvals will be provided within agreed timelines",
            "Client infrastructure meets minimum technical requirements"
        ]
        for assump in standard_assumptions:
            doc.add_paragraph(assump, style="List Bullet")
        
        doc.add_page_break()

        # Section 13 — Acceptance Criteria
        h = doc.add_heading("12. Acceptance & Sign-Off", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        p = doc.add_paragraph(f"Each deliverable will be subject to a formal acceptance process. The Client will have {inputs.change_control_threshold_days} business days to review and accept or reject deliverables with written feedback. Silence after the review period constitutes acceptance.")
        p.runs[0].font.size = Pt(11)
        
        doc.add_page_break()

        # Section 14 — Terms & Conditions
        h = doc.add_heading("13. Terms & Conditions", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        tc_texts = [
            "Confidentiality: Both parties agree to maintain the confidentiality of any proprietary information shared during this engagement.",
            "Limitation of Liability: The total liability of either party shall not exceed the total value of this Statement of Work.",
            "Termination: Either party may terminate this SOW with 30 days written notice. Payment for work completed up to the termination date shall be due.",
            "Governing Law: This Statement of Work shall be governed by the laws of the jurisdiction in which the Service Provider is registered."
        ]
        for tc in tc_texts:
            p = doc.add_paragraph(tc)
            p.runs[0].font.size = Pt(11)
        
        doc.add_page_break()

        # Section 15 — Signature Block
        h = doc.add_heading("14. Signatures", level=1)
        h.runs[0].font.color.rgb = data['primary_rgb']
        
        doc.add_paragraph("By signing below, both parties agree to the terms of this Statement of Work.")
        
        sig_table = doc.add_table(rows=1, cols=2)
        # Col 1 — Service Provider
        c1 = sig_table.cell(0, 0)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run("SIGNED FOR AND ON BEHALF OF:")
        r1.bold = True
        
        p = c1.add_paragraph(data['division'].name)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(12)
        
        c1.add_paragraph("\nSignature: _______________________")
        c1.add_paragraph(f"Name: {data['division_head'].first_name} {data['division_head'].last_name}")
        c1.add_paragraph("Title: Division Head")
        c1.add_paragraph("Date: _______________________")
        
        # Col 2 — Client
        c2 = sig_table.cell(0, 1)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run("SIGNED FOR AND ON BEHALF OF:")
        r2.bold = True
        
        p = c2.add_paragraph(data['lead'].company_name)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(12)
        
        c2.add_paragraph("\nSignature: _______________________")
        c2.add_paragraph("Name: _______________________")
        c2.add_paragraph("Title: _______________________")
        c2.add_paragraph("Date: _______________________")

        # Step 5 — Save to BytesIO
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        file_bytes = output.read()

        # Step 6 — Archive, version, title
        await self.archive_previous_versions(db, lead_id, 'sow')
        version = await self.get_next_version(db, lead_id, 'sow')
        title = self.build_document_title(
            'sow', version, data['lead'].company_name
        )
        filename = f"{title}.docx"

        # Step 7 — Upload to S3
        s3_key = await self.upload_to_s3(
            file_bytes, str(lead_id),
            str(data['division'].id),
            'sow', filename,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        # Step 8 — Upload to SharePoint
        folder = await self.create_sharepoint_folder(
            str(lead_id), data['division'].name,
            data['lead'].company_name, 'sow'
        )
        sharepoint_url = await self.upload_to_sharepoint(
            file_bytes, folder, filename
        )

        # Step 9 — Save document record
        document = await self.save_document_record(
            db=db,
            lead_id=lead_id,
            division_id=data['division'].id,
            template_id=None,
            doc_type='sow',
            format='docx',
            version_number=version,
            title=title,
            s3_key=s3_key,
            sharepoint_url=sharepoint_url,
            created_by=user_id,
            status='pending_manager_review'
        )

        # Step 10 — Notify sales managers for first approval
        manager_stmt = select(User).where(
            User.role == "sales_manager",
            User.division_id == data['division'].id
        )
        result = await db.execute(manager_stmt)
        managers = result.scalars().all()
        
        for manager in managers:
            await create_in_app_notification(
                db=db,
                user_id=manager.id,
                lead_id=lead_id,
                type='doc_generated',
                message=f"SOW v{version} for {data['lead'].company_name} requires your review."
            )

        # Step 11 — Log to activity_timeline
        activity = ActivityTimeline(
            lead_id=lead_id,
            actor_id=user_id,
            event_type=ActivityEventType.document_generated,
            description=f"📑 SOW generated (v{version}) — Pending Manager Review",
            metadata_={
                "doc_type": "sow",
                "format": "docx",
                "version": version,
                "document_id": str(document.id),
                "status": "pending_manager_review",
                "linked_presales_doc_id": str(inputs.linked_presales_doc_id) if inputs.linked_presales_doc_id else None
            }
        )
        db.add(activity)
        await db.commit()

        # Step 12 — Emit socket event
        await socket_manager.emit_to_lead_room(
            lead_id=str(lead_id),
            event="document_generated",
            data={
                "document_id": str(document.id),
                "doc_type": "sow",
                "format": "docx",
                "version": version,
                "title": title,
                "status": "pending_manager_review"
            }
        )

        return document
