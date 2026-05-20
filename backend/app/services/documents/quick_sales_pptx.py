import logging
from io import BytesIO
from datetime import datetime
from uuid import UUID

import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from sqlalchemy import select

from app.models.lead import Lead
from app.models.division import Division
from app.models.user import User
from app.models.activity_timeline import ActivityTimeline, ActivityEventType
from app.models.generated_document import GeneratedDocument
from app.schemas.document import QuickSalesInput
from app.core.socket_manager import socket_manager
from app.services.documents.base_document_service import BaseDocumentService
from app.services.documents.quick_sales_docx import _hex_to_rgb
from app.services.template_service import get_active_template

logger = logging.getLogger(__name__)

def _get_layout(prs, index=6):
    """Return prs.slide_layouts[index]. Use blank layout (index 6) for full design control."""
    return prs.slide_layouts[index]

def _add_text_box(slide, text, left, top, width, height, font_size, bold, color, align):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox

def _add_background(slide, color_rgb):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color_rgb

def _add_logo(slide, logo_url, left, top, width):
    try:
        response = requests.get(logo_url, timeout=5)
        if response.status_code == 200:
            logo_bytes = BytesIO(response.content)
            slide.shapes.add_picture(logo_bytes, left, top, width=width)
    except Exception as e:
        logger.warning(f"Could not load logo for PPTX slide: {e}")

class QuickSalesPptxService(BaseDocumentService):
    SLIDE_WIDTH = Inches(10)
    SLIDE_HEIGHT = Inches(5.625)   # 16:9 ratio

    def _get_layout(self, prs, index=6):
        return _get_layout(prs, index)

    def _add_text_box(self, slide, text, left, top, width, height, font_size, bold, color, align):
        return _add_text_box(slide, text, left, top, width, height, font_size, bold, color, align)

    def _add_background(self, slide, color_rgb):
        return _add_background(slide, color_rgb)

    def _add_logo(self, slide, logo_url, left, top, width):
        return _add_logo(slide, logo_url, left, top, width)

    async def generate(
        self,
        db,
        lead_id: UUID,
        user_id: UUID,
        inputs: QuickSalesInput
    ) -> GeneratedDocument:
        
        # Step 1 — Fetch data from DB
        lead_stmt = select(Lead).where(Lead.id == lead_id)
        result = await db.execute(lead_stmt)
        lead = result.scalar_one()

        division_stmt = select(Division).where(Division.id == lead.division_id)
        result = await db.execute(division_stmt)
        division = result.scalar_one()

        creator_stmt = select(User).where(User.id == user_id)
        result = await db.execute(creator_stmt)
        creator = result.scalar_one()

        branding = division.branding_config or {}
        primary_color = branding.get('primary_color', '#0D1B2A')
        secondary_color = branding.get('secondary_color', '#00C6D7')
        logo_url = branding.get('logo_url', None)
        about_us = branding.get('about_us', '')

        primary_rgb = _hex_to_rgb(primary_color)
        secondary_rgb = _hex_to_rgb(secondary_color)

        # Step 2 — Check for active PPTX template
        template_bytes = await get_active_template(db, division.id, 'quick_sales', 'pptx')
        if template_bytes:
            prs = Presentation(BytesIO(template_bytes))
        else:
            prs = Presentation()
            prs.slide_width = self.SLIDE_WIDTH
            prs.slide_height = self.SLIDE_HEIGHT

        # Step 3 — Build all 8 slides
        
        # --- SLIDE 1 — Cover ---
        slide1 = prs.slides.add_slide(self._get_layout(prs))
        self._add_background(slide1, primary_rgb)
        
        if logo_url:
            self._add_logo(slide1, logo_url, Inches(0.5), Inches(0.3), Inches(1.8))
            
        self._add_text_box(slide1,
            text="Prepared For:",
            left=Inches(1), top=Inches(1.5), width=Inches(8), height=Inches(0.5),
            font_size=14, bold=False, color=RGBColor(255, 255, 255),
            align=PP_ALIGN.CENTER
        )
        self._add_text_box(slide1,
            text=lead.company_name,
            left=Inches(1), top=Inches(2), width=Inches(8), height=Inches(1),
            font_size=32, bold=True, color=secondary_rgb,
            align=PP_ALIGN.CENTER
        )
        self._add_text_box(slide1,
            text=division.name,
            left=Inches(1), top=Inches(3.2), width=Inches(8), height=Inches(0.5),
            font_size=16, bold=False, color=RGBColor(255, 255, 255),
            align=PP_ALIGN.CENTER
        )
        self._add_text_box(slide1,
            text=datetime.today().strftime("%B %d, %Y"),
            left=Inches(1), top=Inches(3.8), width=Inches(8), height=Inches(0.4),
            font_size=12, bold=False, color=RGBColor(200, 200, 200),
            align=PP_ALIGN.CENTER
        )

        # --- SLIDE 2 — About Us ---
        slide2 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide2,
            text="About Us",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=28, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )
        
        line = slide2.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            left=Inches(0.5), top=Inches(1.1),
            width=Inches(9), height=Pt(1)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = secondary_rgb
        line.line.visible = False

        self._add_text_box(slide2,
            text=about_us or "Company overview goes here.",
            left=Inches(0.5), top=Inches(1.3), width=Inches(9), height=Inches(3),
            font_size=14, bold=False, color=RGBColor(50, 50, 50),
            align=PP_ALIGN.LEFT
        )

        stats = [("10+ Years", Inches(0.5)), ("500+ Clients", Inches(3.6)), ("200+ Team", Inches(6.7))]
        for text, left in stats:
            shape = slide2.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                left=left, top=Inches(4.2), width=Inches(2.8), height=Inches(1)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = secondary_rgb
            shape.line.visible = False
            
            self._add_text_box(slide2,
                text=text,
                left=left, top=Inches(4.2), width=Inches(2.8), height=Inches(1),
                font_size=18, bold=True, color=RGBColor(255, 255, 255),
                align=PP_ALIGN.CENTER
            )

        # --- SLIDE 3 — Client Challenges ---
        slide3 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide3,
            text="We Understand Your Challenges",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=24, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )

        for index, challenge in enumerate(inputs.client_challenges[:3]):
            y_position = Inches(1.3) + index * Inches(1.2)
            
            circle = slide3.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                left=Inches(0.5), top=y_position,
                width=Inches(0.6), height=Inches(0.6)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = secondary_rgb
            circle.line.visible = False
            
            self._add_text_box(slide3,
                text=str(index + 1),
                left=Inches(0.5), top=y_position, width=Inches(0.6), height=Inches(0.6),
                font_size=14, bold=True, color=RGBColor(255, 255, 255),
                align=PP_ALIGN.CENTER
            )
            
            self._add_text_box(slide3,
                text=challenge,
                left=Inches(1.3), top=y_position, width=Inches(8.2), height=Inches(0.6),
                font_size=14, bold=False, color=RGBColor(50, 50, 50),
                align=PP_ALIGN.LEFT
            )

        # --- SLIDE 4 — Our Solution ---
        slide4 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide4,
            text="Our Proposed Solution",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=24, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )

        # Left column
        self._add_text_box(slide4,
            text=inputs.proposed_solution_name,
            left=Inches(0.5), top=Inches(1.3), width=Inches(4), height=Inches(0.5),
            font_size=18, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )
        self._add_text_box(slide4,
            text=lead.current_requirement or "Details to be discussed",
            left=Inches(0.5), top=Inches(1.8), width=Inches(4), height=Inches(2.5),
            font_size=13, bold=False, color=RGBColor(80, 80, 80),
            align=PP_ALIGN.LEFT
        )

        # Right column
        self._add_text_box(slide4,
            text="Key Benefits",
            left=Inches(5.2), top=Inches(1.3), width=Inches(4), height=Inches(0.5),
            font_size=16, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )
        benefits_text = "\n".join([f"• {benefit}" for benefit in inputs.key_benefits[:5]])
        self._add_text_box(slide4,
            text=benefits_text,
            left=Inches(5.2), top=Inches(1.8), width=Inches(4.3), height=Inches(2.5),
            font_size=13, bold=False, color=RGBColor(80, 80, 80),
            align=PP_ALIGN.LEFT
        )

        # --- SLIDE 5 — Pricing Summary ---
        slide5 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide5,
            text="Pricing Summary",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=24, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )

        table = slide5.shapes.add_table(4, 3, Inches(0.5), Inches(1.2), Inches(9), Inches(3)).table
        
        headers = ["Item", "Details", "Value"]
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = primary_rgb
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.size = Pt(13)
            p.font.color.rgb = RGBColor(255, 255, 255)

        table.cell(1, 0).text = "Proposed Solution"
        table.cell(1, 1).text = inputs.proposed_solution_name
        table.cell(1, 2).text = ""

        table.cell(2, 0).text = "Estimated Value"
        table.cell(2, 1).text = ""
        table.cell(2, 2).text = f"${lead.estimated_value:,.0f}" if lead.estimated_value else "TBD"

        table.cell(3, 0).text = "Pricing Range"
        table.cell(3, 1).text = inputs.pricing_range
        table.cell(3, 2).text = inputs.pricing_range
        
        # Apply primary_rgb to row 3
        for i in range(3):
            cell = table.cell(3, i)
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = primary_rgb

        # --- SLIDE 6 — Timeline ---
        slide6 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide6,
            text="Project Timeline",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=24, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )

        milestones = [
            ("Kickoff", inputs.start_date, Inches(0.5)),
            ("Implementation", "In Progress", Inches(3.8)),
            ("Delivery", inputs.end_date, Inches(7))
        ]

        for i, (label, date, left) in enumerate(milestones):
            shape = slide6.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                left=left, top=Inches(1.5), width=Inches(2.5), height=Inches(2)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = secondary_rgb
            shape.line.color.rgb = primary_rgb
            
            self._add_text_box(slide6,
                text=label,
                left=left, top=Inches(1.8), width=Inches(2.5), height=Inches(0.5),
                font_size=16, bold=True, color=RGBColor(255, 255, 255),
                align=PP_ALIGN.CENTER
            )
            self._add_text_box(slide6,
                text=date,
                left=left, top=Inches(2.4), width=Inches(2.5), height=Inches(0.4),
                font_size=12, bold=False, color=RGBColor(255, 255, 255),
                align=PP_ALIGN.CENTER
            )

            if i < 2:
                # Add arrow
                arrow_left = left + Inches(2.5) + Inches(0.1)
                slide6.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    arrow_left, Inches(2.5), arrow_left + Inches(0.6), Inches(2.5)
                ).line.color.rgb = secondary_rgb

        # --- SLIDE 7 — Why Us ---
        slide7 = prs.slides.add_slide(self._get_layout(prs))
        self._add_text_box(slide7,
            text="Why Choose Us",
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.7),
            font_size=24, bold=True, color=primary_rgb,
            align=PP_ALIGN.LEFT
        )

        benefits = [
            ("🏆", "Proven Track Record", "10+ years of delivering results"),
            ("⚡", "Fast Turnaround", "Rapid deployment and implementation"),
            ("🤝", "Dedicated Support", "24/7 support throughout the engagement")
        ]

        for index, (icon, title, body) in enumerate(benefits):
            x = Inches(0.5) + index * Inches(3.2)
            
            circle = slide7.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                left=x + Inches(1), top=Inches(1.3), width=Inches(1), height=Inches(1)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = secondary_rgb
            circle.line.visible = False
            
            self._add_text_box(slide7,
                text=icon,
                left=x + Inches(1), top=Inches(1.3), width=Inches(1), height=Inches(1),
                font_size=20, bold=False, color=RGBColor(255, 255, 255),
                align=PP_ALIGN.CENTER
            )
            
            self._add_text_box(slide7,
                text=title,
                left=x, top=Inches(2.5), width=Inches(3), height=Inches(0.5),
                font_size=14, bold=True, color=primary_rgb,
                align=PP_ALIGN.CENTER
            )
            self._add_text_box(slide7,
                text=body,
                left=x, top=Inches(3.1), width=Inches(3), height=Inches(0.8),
                font_size=12, bold=False, color=RGBColor(80, 80, 80),
                align=PP_ALIGN.CENTER
            )

        # --- SLIDE 8 — Contact ---
        slide8 = prs.slides.add_slide(self._get_layout(prs))
        self._add_background(slide8, primary_rgb)
        
        self._add_text_box(slide8,
            text="Let's Get Started",
            left=Inches(1), top=Inches(0.5), width=Inches(8), height=Inches(0.8),
            font_size=28, bold=True, color=RGBColor(255, 255, 255),
            align=PP_ALIGN.CENTER
        )

        card = slide8.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            left=Inches(2.5), top=Inches(1.5), width=Inches(5), height=Inches(3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = secondary_rgb

        self._add_text_box(slide8,
            text=f"{creator.first_name} {creator.last_name}",
            left=Inches(2.5), top=Inches(1.8), width=Inches(5), height=Inches(0.5),
            font_size=18, bold=True, color=primary_rgb,
            align=PP_ALIGN.CENTER
        )
        self._add_text_box(slide8,
            text=creator.email,
            left=Inches(2.5), top=Inches(2.4), width=Inches(5), height=Inches(0.4),
            font_size=13, bold=False, color=RGBColor(80, 80, 80),
            align=PP_ALIGN.CENTER
        )
        self._add_text_box(slide8,
            text=division.name,
            left=Inches(2.5), top=Inches(3.0), width=Inches(5), height=Inches(0.4),
            font_size=12, bold=False, color=secondary_rgb,
            align=PP_ALIGN.CENTER
        )

        if logo_url:
            self._add_logo(slide8, logo_url, Inches(4.3), Inches(4.8), Inches(1.2))

        # Step 4 — Save to BytesIO
        output = BytesIO()
        prs.save(output)
        output.seek(0)
        file_bytes = output.read()

        # Step 5 — Archive previous versions
        await self.archive_previous_versions(db, lead_id, 'quick_sales')

        # Step 6 — Get next version
        version = await self.get_next_version(db, lead_id, 'quick_sales')

        # Step 7 — Build title
        title = self.build_document_title('quick_sales', version, lead.company_name)
        filename = f"{title}.pptx"

        # Step 8 — Upload to S3
        s3_key = await self.upload_to_s3(
            file_bytes, str(lead_id), str(division.id),
            'quick_sales', filename,
            'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )

        # Step 9 — Upload to SharePoint
        folder = await self.create_sharepoint_folder(
            str(lead_id), division.name, lead.company_name, 'quick_sales'
        )
        sharepoint_url = await self.upload_to_sharepoint(file_bytes, folder, filename)

        # Step 10 — Save document record
        document = await self.save_document_record(
            db=db,
            lead_id=lead_id,
            division_id=division.id,
            template_id=None,
            doc_type='quick_sales',
            format='pptx',
            version_number=version,
            title=title,
            s3_key=s3_key,
            sharepoint_url=sharepoint_url,
            created_by=user_id,
            status='draft'
        )

        # Step 11 — Log to activity_timeline
        activity = ActivityTimeline(
            lead_id=lead_id,
            actor_id=user_id,
            event_type=ActivityEventType.document_generated,
            description=f"📊 Quick Sales Presentation generated (v{version})",
            metadata_={
                "doc_type": "quick_sales",
                "format": "pptx",
                "version": version,
                "document_id": str(document.id)
            }
        )
        db.add(activity)
        await db.commit()

        # Step 12 — Emit socket event
        await socket_manager.emit_to_lead_room(
            lead_id=str(lead_id),
            event="document_generated",
            data={
                "document_id": str(document.id),
                "doc_type": "quick_sales",
                "format": "pptx",
                "version": version,
                "title": title
            }
        )

        return document
