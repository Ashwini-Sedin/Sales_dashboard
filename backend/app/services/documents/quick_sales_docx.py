import logging
from io import BytesIO
from datetime import datetime
from uuid import UUID

import requests
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from sqlalchemy import select

from app.models.lead import Lead
from app.models.division import Division
from app.models.user import User
from app.models.activity_timeline import ActivityTimeline, ActivityEventType
from app.models.generated_document import GeneratedDocument
from app.schemas.document import QuickSalesInput
from app.core.socket_manager import socket_manager
from app.services.documents.base_document_service import BaseDocumentService
from app.services.template_service import get_active_template

logger = logging.getLogger(__name__)

def _apply_cell_color(cell, hex_color: str):
    """Apply background color to a table cell"""
    hex_color = hex_color.lstrip('#')
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading)

def _hex_to_rgb(hex_color: str) -> RGBColor:
    """Convert hex color string to RGBColor"""
    hex_color = hex_color.lstrip('#')
    r, g, b = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    return RGBColor(r, g, b)

class QuickSalesDocxService(BaseDocumentService):
    async def generate(
        self,
        db,
        lead_id: UUID,
        user_id: UUID,
        inputs: QuickSalesInput
    ) -> GeneratedDocument:
        
        # Step 1 — Fetch data from DB
        lead_stmt = select(Lead).where(Lead.id == lead_id)
        result = await db.execute(lead_stmt)
        lead = result.scalar_one()

        division_stmt = select(Division).where(Division.id == lead.division_id)
        result = await db.execute(division_stmt)
        division = result.scalar_one()

        creator_stmt = select(User).where(User.id == user_id)
        result = await db.execute(creator_stmt)
        creator = result.scalar_one()

        branding = division.branding_config or {}
        primary_color = branding.get('primary_color', '#0D1B2A')
        secondary_color = branding.get('secondary_color', '#00C6D7')
        logo_url = branding.get('logo_url', None)
        about_us = branding.get('about_us', '')

        # Step 2 — Check for active template
        template_bytes = await get_active_template(db, division.id, 'quick_sales', 'docx')
        if template_bytes:
            doc = Document(BytesIO(template_bytes))
        else:
            doc = Document()

        # Step 3 — Build the document sections
        
        # Section 1 — Cover Page
        p = doc.add_paragraph()
        run = p.add_run(division.name)
        run.font.size = Pt(28)
        run.bold = True
        run.font.color.rgb = _hex_to_rgb(primary_color)
        
        doc.add_paragraph("Prepared For:").runs[0].font.size = Pt(14)
        
        p = doc.add_paragraph()
        run = p.add_run(lead.company_name)
        run.bold = True
        run.font.size = Pt(20)
        
        doc.add_paragraph(f"{lead.first_name} {lead.last_name}").runs[0].font.size = Pt(14)
        doc.add_paragraph(f"Prepared By: {division.name}").runs[0].font.size = Pt(12)
        doc.add_paragraph(datetime.today().strftime("%B %d, %Y")).runs[0].font.size = Pt(12)

        if logo_url:
            try:
                response = requests.get(logo_url)
                if response.status_code == 200:
                    doc.add_picture(BytesIO(response.content), width=Inches(2.0))
            except Exception as e:
                logger.warning(f"Could not load logo, skipping: {e}")

        doc.add_page_break()

        # Section 2 — About Us
        heading = doc.add_heading("About Us", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        doc.add_paragraph(about_us if about_us else "Company overview goes here.")
        doc.add_page_break()

        # Section 3 — Understanding Your Challenges
        heading = doc.add_heading("Understanding Your Challenges", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        for challenge in inputs.client_challenges:
            doc.add_paragraph(challenge, style='List Bullet')
        doc.add_page_break()

        # Section 4 — Our Proposed Solution
        heading = doc.add_heading("Our Proposed Solution", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        
        p = doc.add_paragraph()
        run = p.add_run(inputs.proposed_solution_name)
        run.bold = True
        run.font.size = Pt(14)
        
        doc.add_paragraph(lead.current_requirement or "Solution details to be discussed.")
        doc.add_paragraph("Key Benefits:").runs[0].bold = True
        for benefit in inputs.key_benefits:
            doc.add_paragraph(benefit, style='List Bullet')
        doc.add_page_break()

        # Section 4b — Professional Information Note
        if inputs.introduction or inputs.objective or inputs.content_structure:
            heading = doc.add_heading("Professional Information Note", level=1)
            heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
            
            if inputs.introduction:
                p = doc.add_paragraph()
                run = p.add_run(inputs.introduction)
                run.font.size = Pt(11)
            
            if inputs.objective:
                p = doc.add_heading("Objective", level=2)
                p.runs[0].font.color.rgb = _hex_to_rgb(secondary_color)
                p = doc.add_paragraph()
                run = p.add_run(inputs.objective)
                run.font.size = Pt(11)
                
            if inputs.content_structure:
                p = doc.add_heading("Content Structure", level=2)
                p.runs[0].font.color.rgb = _hex_to_rgb(secondary_color)
                p = doc.add_paragraph()
                run = p.add_run(inputs.content_structure)
                run.font.size = Pt(11)
                
            doc.add_page_break()

        # Section 5 — Pricing Summary
        heading = doc.add_heading("Pricing Summary", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        for i, text in enumerate(['Item', 'Details', 'Value']):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _apply_cell_color(hdr_cells[i], primary_color)

        row_cells = table.add_row().cells
        row_cells[0].text = "Proposed Solution"
        row_cells[1].text = inputs.proposed_solution_name
        row_cells[2].text = ""

        row_cells = table.add_row().cells
        row_cells[0].text = "Estimated Value"
        row_cells[1].text = ""
        row_cells[2].text = f"${lead.estimated_value:,.0f}" if lead.estimated_value else "TBD"

        row_cells = table.add_row().cells
        row_cells[0].text = "Pricing Range"
        row_cells[1].text = inputs.pricing_range
        row_cells[2].text = ""

        row_cells = table.add_row().cells
        row_cells[0].text = "Investment Range"
        row_cells[1].text = ""
        row_cells[2].text = inputs.pricing_range
        for i in range(3):
            row_cells[i].paragraphs[0].runs[0].bold = True
            row_cells[i].paragraphs[0].runs[0].font.color.rgb = _hex_to_rgb(primary_color)
            
        doc.add_page_break()

        # Section 6 — Project Timeline
        heading = doc.add_heading("Project Timeline", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        for i, text in enumerate(['Phase', 'Start Date', 'End Date']):
            hdr_cells[i].text = text
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            _apply_cell_color(hdr_cells[i], primary_color)
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        row_cells = table.add_row().cells
        row_cells[0].text = "Project Kickoff"
        row_cells[1].text = inputs.start_date
        row_cells[2].text = ""

        row_cells = table.add_row().cells
        row_cells[0].text = "Implementation"
        row_cells[1].text = ""
        row_cells[2].text = ""

        row_cells = table.add_row().cells
        row_cells[0].text = "Delivery"
        row_cells[1].text = ""
        row_cells[2].text = inputs.end_date
        
        doc.add_page_break()

        # Section 7 — Next Steps & Contact
        heading = doc.add_heading("Next Steps & Contact", level=1)
        heading.runs[0].font.color.rgb = _hex_to_rgb(primary_color)
        
        for step in ["Schedule a follow-up meeting", "Review and finalize proposal", "Sign engagement agreement"]:
            doc.add_paragraph(step, style='List Number')
            
        doc.add_paragraph("")
        doc.add_paragraph("Your Contact:").runs[0].bold = True
        doc.add_paragraph(f"{creator.first_name} {creator.last_name}")
        doc.add_paragraph(creator.email)
        if creator.bio:
            doc.add_paragraph(creator.bio).runs[0].italic = True

        # Step 4 — Save to BytesIO
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        file_bytes = output.read()

        # Step 5 — Archive previous versions
        await self.archive_previous_versions(db, lead_id, 'quick_sales')

        # Step 6 — Get next version
        version = await self.get_next_version(db, lead_id, 'quick_sales')

        # Step 7 — Build title
        title = self.build_document_title('quick_sales', version, lead.company_name)
        filename = f"{title}.docx"

        # Step 8 — Upload to S3
        s3_key = await self.upload_to_s3(
            file_bytes, str(lead_id), str(division.id),
            'quick_sales', filename,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        # Step 9 — Upload to SharePoint
        folder = await self.create_sharepoint_folder(
            str(lead_id), division.name, lead.company_name, 'quick_sales'
        )
        sharepoint_url = await self.upload_to_sharepoint(file_bytes, folder, filename)

        # Step 10 — Save document record
        document = await self.save_document_record(
            db=db,
            lead_id=lead_id,
            division_id=division.id,
            template_id=None,
            doc_type='quick_sales',
            format='docx',
            version_number=version,
            title=title,
            s3_key=s3_key,
            sharepoint_url=sharepoint_url,
            created_by=user_id,
            status='draft'
        )

        # Step 11 — Log to activity_timeline
        activity = ActivityTimeline(
            lead_id=lead_id,
            actor_id=user_id,
            event_type=ActivityEventType.document_generated,
            description=f"📄 Quick Sales Document generated (v{version})",
            metadata_={
                "doc_type": "quick_sales",
                "format": "docx",
                "version": version,
                "document_id": str(document.id)
            }
        )
        db.add(activity)
        await db.commit()

        # Step 12 — Emit socket event
        await socket_manager.emit_to_lead_room(
            lead_id=str(lead_id),
            event="document_generated",
            data={
                "document_id": str(document.id),
                "doc_type": "quick_sales",
                "format": "docx",
                "version": version,
                "title": title
            }
        )

        return document
