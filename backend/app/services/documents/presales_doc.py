import logging
from datetime import datetime
from io import BytesIO
from uuid import UUID
from itertools import groupby

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.models.lead import Lead
from app.models.division import Division
from app.models.user import User
from app.models.tech_library import TechLibrary
from app.models.generated_document import GeneratedDocument
from app.models.activity_timeline import ActivityTimeline, ActivityEventType
from app.schemas.document import PresalesInput
from app.services.documents.base_document_service import BaseDocumentService
from app.services.template_service import get_active_template
from app.services.notification_service import create_in_app_notification
from app.core.socket_manager import socket_manager
from sqlalchemy import select

logger = logging.getLogger(__name__)

class PresalesDocService(BaseDocumentService):
    async def generate_docx(
        self,
        db,
        lead_id: UUID,
        user_id: UUID,
        inputs: PresalesInput
    ) -> GeneratedDocument:
        # Step 1: Fetch data from DB
        lead = (await db.execute(select(Lead).where(Lead.id == lead_id))).scalar_one()
        division = (await db.execute(select(Division).where(Division.id == lead.division_id))).scalar_one()
        creator = (await db.execute(select(User).where(User.id == user_id))).scalar_one()
        
        # Fetch division head (assuming division has head_user_id)
        # Note: If Division model doesn't have head_user_id, we might need to adjust.
        # But instructions say "fetch User where id = division.head_user_id"
        division_head = None
        if hasattr(division, 'head_user_id') and division.head_user_id:
            division_head = (await db.execute(select(User).where(User.id == division.head_user_id))).scalar_one_or_none()
        
        technologies = (await db.execute(
            select(TechLibrary).where(
                TechLibrary.id.in_(inputs.technology_ids),
                TechLibrary.is_active == True
            )
        )).scalars().all()

        branding = division.branding_config or {}
        primary_color = branding.get('primary_color', '#0D1B2A')
        secondary_color = branding.get('secondary_color', '#00C6D7')
        logo_url = branding.get('logo_url', None)

        primary_rgb = self._hex_to_rgb(primary_color)
        secondary_rgb = self._hex_to_rgb(secondary_color)

        # Step 2: Check for active template
        template_bytes = await get_active_template(db, division.id, 'presales', 'docx')
        if template_bytes:
            doc = Document(BytesIO(template_bytes))
        else:
            doc = Document()

        # Step 3: Calculate effort totals
        total_days = sum(row.days for row in inputs.effort_rows)
        total_cost = sum(row.days * row.rate_per_day for row in inputs.effort_rows)

        # Step 4: Build document sections
        
        # Section 1: Cover Page
        self._add_cover_page(doc, division.name, lead, inputs.commercial_model, primary_rgb, secondary_rgb, logo_url)
        doc.add_page_break()

        # Section 2: Executive Overview
        self._add_heading(doc, "Executive Overview", primary_rgb)
        doc.add_paragraph(inputs.executive_overview)
        doc.add_page_break()

        # Section 3: Client Background
        self._add_heading(doc, "Client Background", primary_rgb)
        doc.add_paragraph(inputs.client_background)
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        data = [
            ("Company", lead.company_name),
            ("Contact", f"{lead.first_name} {lead.last_name}"),
            ("Email", lead.email),
            ("Commercial Model", inputs.commercial_model.title())
        ]
        for i, (label, value) in enumerate(data):
            table.cell(i, 0).text = label
            table.cell(i, 1).text = value
        doc.add_page_break()

        # Section 4: Problem Statement
        self._add_heading(doc, "Problem Statement", primary_rgb)
        doc.add_paragraph(inputs.problem_statement)
        doc.add_page_break()

        # Section 5: Solution Overview
        self._add_heading(doc, "Solution Overview", primary_rgb)
        doc.add_paragraph(inputs.solution_overview)
        doc.add_page_break()

        # Section 6: Scope of Work
        self._add_heading(doc, "Scope of Work", primary_rgb)
        in_scope_items = [i.item for i in inputs.scope_items if i.in_scope]
        out_of_scope_items = [i.item for i in inputs.scope_items if not i.in_scope]
        
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'In Scope'
        hdr_cells[1].text = 'Out of Scope'
        self._set_cell_background(hdr_cells[0], primary_color)
        self._set_cell_background(hdr_cells[1], primary_color)
        
        max_rows = max(len(in_scope_items), len(out_of_scope_items))
        for idx in range(max_rows):
            row_cells = table.add_row().cells
            if idx < len(in_scope_items):
                row_cells[0].text = "✅ " + in_scope_items[idx]
            if idx < len(out_of_scope_items):
                row_cells[1].text = "❌ " + out_of_scope_items[idx]
        doc.add_page_break()

        # Section 7: Assumptions
        self._add_heading(doc, "Assumptions", primary_rgb)
        for assumption in inputs.assumptions:
            doc.add_paragraph(assumption, style='List Bullet')
        doc.add_page_break()

        # Section 8: Effort Estimate
        self._add_heading(doc, "Effort Estimate", primary_rgb)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        for i, text in enumerate(['Phase', 'Role', 'Days', 'Rate/Day']):
            hdr_cells[i].text = text
            self._set_cell_background(hdr_cells[i], primary_color)
            
        for row in inputs.effort_rows:
            row_cells = table.add_row().cells
            row_cells[0].text = row.phase
            row_cells[1].text = row.role
            row_cells[2].text = str(row.days)
            row_cells[3].text = f"${row.rate_per_day:,.0f}"
            
        total_row = table.add_row().cells
        total_row[0].text = "TOTAL"
        total_row[2].text = f"{total_days} days"
        total_row[3].text = f"${total_cost:,.0f}"
        for cell in total_row:
            self._set_cell_background(cell, primary_color)
            
        note = doc.add_paragraph(f"Commercial Model: {inputs.commercial_model.title()}")
        note.runs[0].italic = True
        doc.add_page_break()

        # Section 9: Risk Register
        self._add_heading(doc, "Risk Register", primary_rgb)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        for i, text in enumerate(['Risk', 'Impact', 'Probability', 'Mitigation']):
            hdr_cells[i].text = text
            self._set_cell_background(hdr_cells[i], primary_color)
            
        for risk in inputs.risks:
            row_cells = table.add_row().cells
            row_cells[0].text = risk.risk
            row_cells[1].text = risk.impact
            row_cells[2].text = risk.probability
            row_cells[3].text = risk.mitigation
            
            # Apply color to impact cell
            color_map = {
                "High": ("FFE0E0", RGBColor(0xFF, 0x00, 0x00)),
                "Medium": ("FFF0D0", RGBColor(0xFF, 0xA5, 0x00)),
                "Low": ("E0FFE0", RGBColor(0x00, 0xB0, 0x50))
            }
            if risk.impact in color_map:
                bg, fg = color_map[risk.impact]
                self._set_cell_background(row_cells[1], bg)
                row_cells[1].paragraphs[0].runs[0].font.color.rgb = fg
        doc.add_page_break()

        # Section 10: Technology Stack
        self._add_heading(doc, "Technology Stack", primary_rgb)
        sorted_techs = sorted(technologies, key=lambda t: t.category)
        for category, techs in groupby(sorted_techs, key=lambda t: t.category):
            p = doc.add_paragraph()
            run = p.add_run(category.title())
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = secondary_rgb
            for tech in techs:
                doc.add_paragraph(tech.name, style='List Bullet')
        doc.add_page_break()

        # Section 11: Approval & Sign-Off
        self._add_heading(doc, "Approval & Sign-Off", primary_rgb)
        doc.add_paragraph("This document requires approval from the Division Head before proceeding.")
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Name"
        hdr_cells[1].text = "Role"
        hdr_cells[2].text = "Signature"
        for cell in hdr_cells:
            self._set_cell_background(cell, primary_color)
            
        if division_head:
            row1 = table.rows[1].cells
            row1[0].text = f"{division_head.first_name} {division_head.last_name}"
            row1[1].text = "Division Head"
            
        row2 = table.rows[2].cells
        row2[0].text = f"{creator.first_name} {creator.last_name}"
        row2[1].text = creator.role.replace("_", " ").title()
        
        status_note = doc.add_paragraph("Status: Pending Approval")
        status_note.runs[0].italic = True

        # Step 5: Save to BytesIO
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        file_bytes = output.read()

        # Step 6: Archive previous versions
        await self.archive_previous_versions(db, lead_id, 'presales')

        # Step 7: Get next version
        version = await self.get_next_version(db, lead_id, 'presales')

        # Step 8: Build title and filename
        title = self.build_document_title('presales', version, lead.company_name)
        filename = f"{title}.docx"

        # Step 9: Upload to S3
        s3_key = await self.upload_to_s3(
            file_bytes, str(lead_id), str(division.id),
            'presales', filename,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        # Step 10: Upload to SharePoint
        folder = await self.create_sharepoint_folder(
            str(lead_id), division.name, lead.company_name, 'presales'
        )
        sharepoint_url = await self.upload_to_sharepoint(file_bytes, folder, filename)

        # Step 11: Save document record
        document = await self.save_document_record(
            db=db,
            lead_id=lead_id,
            division_id=division.id,
            template_id=None,
            doc_type='presales',
            format='docx',
            version_number=version,
            title=title,
            s3_key=s3_key,
            sharepoint_url=sharepoint_url,
            created_by=user_id,
            status='pending_approval'
        )

        # Step 12: Notify division head
        if division_head:
            await create_in_app_notification(
                db=db,
                user_id=division_head.id,
                lead_id=lead_id,
                type='doc_generated',
                message=f"Presales Document v{version} for {lead.company_name} requires your approval."
            )

        # Step 13: Log to activity_timeline
        activity = ActivityTimeline(
            lead_id=lead_id,
            actor_id=user_id,
            event_type=ActivityEventType.document_generated,
            description=f"📋 Presales Document generated (v{version}) — Pending Approval",
            metadata_={
                "doc_type": "presales",
                "format": "docx",
                "version": version,
                "document_id": str(document.id),
                "status": "pending_approval"
            }
        )
        db.add(activity)
        await db.commit()

        # Step 14: Emit socket event
        await socket_manager.emit_to_lead_room(
            lead_id=str(lead_id),
            event="document_generated",
            data={
                "document_id": str(document.id),
                "doc_type": "presales",
                "format": "docx",
                "version": version,
                "title": title,
                "status": "pending_approval"
            }
        )

        return document

    def _hex_to_rgb(self, hex_color):
        hex_color = hex_color.lstrip('#')
        return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))

    def _add_cover_page(self, doc, division_name, lead, comm_model, primary_rgb, secondary_rgb, logo_url):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(division_name)
        run.bold = True
        run.font.size = Pt(28)
        run.font.color.rgb = primary_rgb

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Presales Document")
        run.font.size = Pt(20)
        run.font.color.rgb = secondary_rgb

        doc.add_paragraph("\n" * 2)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Prepared For: {lead.company_name}")
        run.font.size = Pt(16)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Contact: {lead.first_name} {lead.last_name}")
        run.font.size = Pt(14)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Commercial Model: {comm_model.upper()}")
        run.font.size = Pt(12)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Date: {datetime.today().strftime('%B %d, %Y')}")
        run.font.size = Pt(12)

        if logo_url:
            # Adding a placeholder for logo as we can't fetch it easily from URL here without requests
            pass

    def _add_heading(self, doc, text, color_rgb):
        h = doc.add_heading(text, level=1)
        h.runs[0].font.color.rgb = color_rgb

    def _set_cell_background(self, cell, hex_color):
        if hex_color.startswith('#'):
            hex_color = hex_color[1:]
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)
